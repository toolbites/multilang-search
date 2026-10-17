"""
LinkFinder backend
English keyword -> har language mein translate -> us mulk ke search results (kai pages)
-> duplicates hata kar -> website traffic -> JSON / Excel / CSV

Commands:
  run     Sab kuch ek saath (apne computer par)
            python search_tool.py run -k "free games" -l all -p 2
  search  Sirf languages ka ek hissa search kare (GitHub parallel jobs)
            python search_tool.py search -k "free games" -l de fr ja -p 2 -o part_1.json
  merge   Saare hisse jor kar duplicates hataye + traffic
            python search_tool.py merge -i parts/ -o result.json --xlsx out.xlsx
  chunks  Languages ko N hisson mein baante (GitHub matrix ke liye)

Engines: free (default, bina key)  |  serper (asli Google, SERPER_API_KEY)
"""
import argparse
import csv
import glob
import json
import os
import random
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlparse, urlunparse, parse_qsl, urlencode

import requests
from deep_translator import GoogleTranslator

from languages import LANGUAGES, resolve
from secure import decrypt, encrypt, mask

PER_PAGE = 10
TRACKING = {"utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content", "gclid", "fbclid", "ref"}


# ---------------- helpers ----------------
def log(*a):
    print(*a, flush=True)


class TranslateError(Exception):
    pass


def _gtx(text, code):
    """Google Translate (gtx) – tez, JSON."""
    r = requests.get("https://translate.googleapis.com/translate_a/single",
                     params={"client": "gtx", "sl": "en", "tl": code, "dt": "t", "q": text}, timeout=15,
                     headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code == 429:
        raise TranslateError("429")
    r.raise_for_status()
    out = "".join(part[0] for part in r.json()[0] if part and part[0])
    return out.strip()


def _deep(text, code):
    return (GoogleTranslator(source="en", target=code).translate(text) or "").strip()


def _mymemory(text, code):
    tl = {"iw": "he", "jw": "jv", "zh-CN": "zh-CN", "zh-TW": "zh-TW", "mni-Mtei": "mni"}.get(code, code)
    r = requests.get("https://api.mymemory.translated.net/get",
                     params={"q": text, "langpair": f"en|{tl}"}, timeout=15)
    r.raise_for_status()
    d = r.json()
    if int(d.get("responseStatus", 200)) != 200:
        raise TranslateError(str(d.get("responseDetails", ""))[:80])
    return (d.get("responseData", {}).get("translatedText") or "").strip()


def translate(text, code):
    """Translation ya None (fail). Kabhi chup chaap English wapas nahi deta."""
    if code == "en":
        return text
    last = "empty response"
    for attempt in range(4):
        for fn in (_gtx, _deep, _mymemory):
            try:
                out = fn(text, code)
                if out and "MYMEMORY WARNING" not in out.upper():
                    return out
            except Exception as e:
                last = f"{fn.__name__}: {e}"
        time.sleep(2 + attempt * 3 + random.random() * 2)
    log(f"  ! translate fail {code}: {last[:120]}")
    return None


def translate_all(keyword, codes):
    """Plan step mein ek baar, aaram se. {code: translation}"""
    out = {}
    for i, c in enumerate(codes):
        t = translate(keyword, LANGUAGES[c]["translate"])
        if t:
            out[c] = t
            mask(t)
        time.sleep(0.15 + random.random() * 0.2)
    log(f"Translated {len(out)}/{len(codes)} languages")
    return out


def norm_url(u):
    """Duplicate pakadne ke liye URL saaf: http/https, www, trailing /, tracking params."""
    try:
        p = urlparse(u.strip())
        host = p.netloc.lower().removeprefix("www.")
        q = urlencode([(k, v) for k, v in parse_qsl(p.query) if k.lower() not in TRACKING])
        path = p.path.rstrip("/") or ""
        return urlunparse(("", host, path, "", q, ""))
    except Exception:
        return u


def domain_of(u):
    return urlparse(u).netloc.lower().split(":")[0].removeprefix("www.")


# ---------------- engines ----------------
# Wikipedia/Grokipedia nahi: sirf asli search engines, is order mein try
FREE_ENGINES = ["duckduckgo", "brave", "bing", "yahoo", "mojeek", "google", "yandex", "startpage"]
SKIP_DOMAINS = ("wikipedia.org", "grokipedia.com")
MIN_GOOD = 5  # itne results mile to engine theek maana jaye


def _free_page(ddg, query, region, page, engine):
    res = ddg.text(query, region=region, safesearch="moderate", max_results=PER_PAGE, page=page, backend=engine) or []
    out = []
    for r in res:
        u = r.get("href") or r.get("url") or ""
        if not u.startswith("http") or any(d in u for d in SKIP_DOMAINS):
            continue
        out.append({"title": r.get("title", ""), "url": u, "snippet": r.get("body", "")})
    return out


def search_free(query, lang, pages):
    """Har page alag se. Pehla engine jo achhe results de, baqi pages bhi usi se (rank sahi rahe)."""
    from ddgs import DDGS
    ddg = DDGS(timeout=12)
    region = lang["region"]
    out, chosen, stats = [], None, []
    for page in range(1, pages + 1):
        engines = [chosen] if chosen else FREE_ENGINES
        best, best_eng = [], None
        for eng in engines:
            try:
                got = _free_page(ddg, query, region, page, eng)
            except Exception:
                stats.append(f"{eng}:err")
                time.sleep(0.4 + random.random() * 0.6)
                continue
            if len(got) > len(best):
                best, best_eng = got, eng
            if len(got) >= (MIN_GOOD if not chosen else 1):
                break
            stats.append(f"{eng}:{len(got)}")
            time.sleep(0.4 + random.random() * 0.6)
        if not best:
            break  # aage ke pages bhi nahi milenge
        chosen = chosen or best_eng
        for i, r in enumerate(best[:PER_PAGE]):
            # Rank poori list mein (page 2 ka pehla = 11), Google ki tarah
            out.append({"position": (page - 1) * PER_PAGE + i + 1, "page": page, "engine": best_eng, **r})
        time.sleep(0.5 + random.random())
    tried = ",".join(stats[:6])
    log(f"  engine={chosen or 'none'} results={len(out)}" + (f" (tried {tried})" if tried else ""))
    return out


def search_serper(query, lang, pages, key):
    out, pos = [], 0
    for page in range(1, pages + 1):
        try:
            r = requests.post("https://google.serper.dev/search", timeout=30,
                              headers={"X-API-KEY": key, "Content-Type": "application/json"},
                              json={"q": query, "gl": lang["gl"], "hl": lang["hl"], "page": page, "num": PER_PAGE})
            r.raise_for_status()
            items = r.json().get("organic", [])
        except Exception as e:
            log(f"  ! serper page {page}: {e}")
            break
        if not items:
            break
        for it in items:
            pos += 1
            out.append({"position": pos, "page": page, "title": it.get("title", ""),
                        "url": it.get("link", ""), "snippet": it.get("snippet", "")})
    return out


def one_language(keyword, code, pages, engine, key, tmap=None):
    lang = LANGUAGES[code]
    q = (tmap or {}).get(code) if tmap is not None else translate(keyword, lang["translate"])
    if not q:
        log(f"[{lang['name']}] skipped (no translation)")
        return code, []
    mask(q)
    res = search_serper(q, lang, pages, key) if engine == "serper" else search_free(q, lang, pages)
    rows = []
    for r in res:
        if not r["url"].startswith("http"):
            continue
        rows.append({"language": lang["name"], "lang_code": code, "keyword": keyword,
                     "translated_keyword": q, "google_domain": lang["google_domain"],
                     "domain": domain_of(r["url"]), **r})
    log(f"[{lang['name']}] {len(rows)} links")
    return code, rows


def do_search(keyword, codes, pages, engine, workers, tmap=None):
    key = os.environ.get("SERPER_API_KEY", "")
    if engine == "serper" and not key:
        sys.exit("SERPER_API_KEY set nahi hai.")
    results = {}
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = []
        for i, c in enumerate(codes):
            futs.append(ex.submit(one_language, keyword, c, pages, engine, key, tmap))
            if engine == "free":
                time.sleep(0.8)  # rate limit se bachne ke liye halka gap
        for f in as_completed(futs):
            try:
                c, rows = f.result()
                results[c] = rows
            except Exception as e:
                log(f"  ! language fail: {e}")
    # Asli language order mein wapas
    return [r for c in codes for r in results.get(c, [])]


# ---------------- merge / dedupe ----------------
def dedupe(rows):
    """Same URL sirf ek baar. Jo pehle mila (behtar rank) woh rehta hai;
    baqi languages 'also_in' mein likh di jati hain."""
    rows = sorted(rows, key=lambda r: (r["position"], r["language"]))
    seen = {}
    for r in rows:
        k = norm_url(r["url"])
        if k in seen:
            first = seen[k]
            if r["language"] not in first["also_in"] and r["language"] != first["language"]:
                first["also_in"].append(r["language"])
            continue
        r["also_in"] = []
        seen[k] = r
    out = list(seen.values())
    for r in out:
        r["found_in_languages"] = 1 + len(r["also_in"])
        r["also_in"] = ", ".join(r["also_in"])
    # Dikhane ka order: language, phir rank
    order = {c: i for i, c in enumerate(LANGUAGES)}
    out.sort(key=lambda r: (order.get(r["lang_code"], 999), r["position"]))
    return out


def finalize(rows, add_traffic_data=True):
    total = len(rows)
    rows = dedupe(rows)
    if add_traffic_data:
        from traffic import add_traffic
        add_traffic(rows)
    langs = sorted({r["language"] for r in rows})
    return {
        "status": "done",
        "count": len(rows),
        "raw_count": total,
        "duplicates_removed": total - len(rows),
        "unique_websites": len({r["domain"] for r in rows}),
        "languages_searched": len(langs),
        "results": rows,
    }


COLS = ["language", "translated_keyword", "google_domain", "page", "position", "title", "url",
        "domain", "traffic", "monthly_visits", "tranco_rank", "traffic_tier", "traffic_source",
        "found_in_languages", "also_in"]


def write_tables(data, xlsx=None, csv_path=None):
    rows = data["results"]
    if csv_path:
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    if xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "All Links"
        ws.append(COLS)
        for r in rows:
            ws.append([r.get(c, "") for c in COLS])
            cell = ws.cell(ws.max_row, COLS.index("url") + 1)
            cell.hyperlink = r["url"]
            cell.font = Font(color="534AB7", underline="single")
        ws2 = wb.create_sheet("Unique Websites")
        hdr = ["domain", "traffic", "monthly_visits", "tranco_rank", "links_found", "languages", "best_position"]
        ws2.append(hdr)
        agg = {}
        for r in rows:
            a = agg.setdefault(r["domain"], {"r": r, "n": 0, "langs": set(), "best": 999})
            a["n"] += 1
            a["langs"].add(r["language"])
            a["best"] = min(a["best"], r["position"])
        for d, a in sorted(agg.items(), key=lambda x: (x[1]["r"].get("tranco_rank") or 10**9)):
            ws2.append([d, a["r"].get("traffic", ""), a["r"].get("monthly_visits"), a["r"].get("tranco_rank"),
                        a["n"], ", ".join(sorted(a["langs"])), a["best"]])
        for sh in (ws, ws2):
            for c in sh[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="534AB7")
            sh.freeze_panes = "A2"
            sh.auto_filter.ref = sh.dimensions
            for col in sh.columns:
                sh.column_dimensions[col[0].column_letter].width = 20
        wb.save(xlsx)


# ---------------- CLI ----------------
def main():
    ap = argparse.ArgumentParser(description="LinkFinder backend")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p):
        p.add_argument("-k", "--keyword", help="English keyword (ya encrypted LF_PAYLOAD env)")
        p.add_argument("-l", "--languages", nargs="*", default=["all"])
        p.add_argument("-p", "--pages", type=int, default=2)
        p.add_argument("-e", "--engine", choices=["free", "serper"],
                       default="serper" if os.environ.get("SERPER_API_KEY") else "free")
        p.add_argument("-w", "--workers", type=int, default=3)

    p_run = sub.add_parser("run"); common(p_run)
    p_run.add_argument("-o", "--out", default="output")
    p_run.add_argument("--no-traffic", action="store_true")

    p_s = sub.add_parser("search"); common(p_s)
    p_s.add_argument("-o", "--out", required=True)
    p_s.add_argument("--encrypt", action="store_true", help="part file encrypt karein (public repo)")
    p_s.add_argument("--translations", help="translate command ki (encrypted) file")

    p_m = sub.add_parser("merge")
    p_m.add_argument("-i", "--input", required=True, help="parts folder")
    p_m.add_argument("-o", "--out", required=True)
    p_m.add_argument("--xlsx"); p_m.add_argument("--csv")
    p_m.add_argument("--no-traffic", action="store_true")

    p_t = sub.add_parser("translate")
    p_t.add_argument("-o", "--out", required=True)

    p_c = sub.add_parser("chunks")
    p_c.add_argument("-l", "--languages", nargs="*", default=None)
    p_c.add_argument("-j", "--jobs", type=int, default=5, help="zyada se zyada kitne parallel servers")
    p_c.add_argument("--min-size", type=int, default=6, help="har server par kam se kam languages")

    a = ap.parse_args()

    # Website se aaya encrypted payload (keyword/languages/pages)
    payload = None
    if os.environ.get("LF_PAYLOAD"):
        payload = decrypt(os.environ["LF_PAYLOAD"])
        mask(payload.get("keyword", ""))
        if a.cmd in ("search", "run"):
            a.keyword = a.keyword or payload["keyword"]
            a.pages = int(payload.get("pages", a.pages))
        if a.cmd in ("chunks", "translate") and getattr(a, "languages", None) is None:
            a.languages = str(payload.get("languages", "all")).split()
    if a.cmd == "translate":
        codes, _ = resolve(str(payload.get("languages", "all")).split() if payload else ["all"])
        tm = translate_all(payload["keyword"], codes)
        json.dump({"enc": encrypt(tm)}, open(a.out, "w"))
        return

    if a.cmd == "chunks" and not a.languages:
        a.languages = ["all"]
    if a.cmd in ("search", "run") and not a.keyword:
        sys.exit("Keyword nahi mila (-k ya LF_PAYLOAD).")

    if a.cmd == "chunks":
        codes, bad = resolve(a.languages)
        if bad:
            print(f"! ignored codes: {bad}", file=sys.stderr)
        size = max(a.min_size, -(-len(codes) // max(1, a.jobs)))
        parts = [" ".join(codes[i:i + size]) for i in range(0, len(codes), size)] or ["en"]
        print(json.dumps(parts))
        return

    if a.cmd == "merge":
        rows = []
        for f in sorted(glob.glob(os.path.join(a.input, "**", "*.json"), recursive=True)):
            try:
                part = json.load(open(f, encoding="utf-8"))
                if "enc" in part:
                    part = decrypt(part["enc"])
                rows += part.get("results", [])
            except Exception as e:
                log(f"! bad part {f}: {e}")
        data = finalize(rows, not a.no_traffic) if rows else {"status": "failed", "error": "No results found. Try another keyword.", "results": []}
        json.dump(data, open(a.out, "w", encoding="utf-8"), ensure_ascii=False)
        if rows:
            write_tables(data, a.xlsx, a.csv)
        log(f"Merged: {data.get('count', 0)} links ({data.get('duplicates_removed', 0)} duplicates removed)")
        return

    codes, bad = resolve(a.languages)
    if bad:
        sys.exit(f"Ghalat language codes: {bad}")
    pages = max(1, min(10, a.pages))
    tmap = None
    if getattr(a, "translations", None) and os.path.exists(a.translations):
        d = json.load(open(a.translations))
        tmap = decrypt(d["enc"]) if "enc" in d else d
        for v in tmap.values():
            mask(v)
    rows = do_search(a.keyword, codes, pages, a.engine, a.workers, tmap)

    if a.cmd == "search":
        part = {"results": rows}
        if a.encrypt:
            part = {"enc": encrypt(part)}
        json.dump(part, open(a.out, "w", encoding="utf-8"), ensure_ascii=False)
        log(f"Part saved: {len(rows)} links -> {a.out}")
        return

    # run
    os.makedirs(a.out, exist_ok=True)
    data = finalize(rows, not a.no_traffic)
    stamp = time.strftime("%Y%m%d_%H%M")
    base = os.path.join(a.out, f"linkfinder_{stamp}")
    json.dump(data, open(base + ".json", "w", encoding="utf-8"), ensure_ascii=False)
    write_tables(data, base + ".xlsx", base + ".csv")
    log(f"\nDone: {data['count']} links, {data['duplicates_removed']} duplicates removed, "
        f"{data['unique_websites']} websites\n  {base}.xlsx")


if __name__ == "__main__":
    main()
